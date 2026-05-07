"""Excel report generation for network pharmacology analysis."""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, BubbleChart, Series
from openpyxl.drawing.image import Image
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import io
from datetime import datetime
from typing import Optional


TCM_GREEN = "1E5631"
TCM_LIGHT = "D5E8D4"
TCM_ACCENT = "82B366"
HEADER_FILL = PatternFill(start_color=TCM_GREEN, end_color=TCM_GREEN, fill_type="solid")
SUB_FILL = PatternFill(start_color=TCM_LIGHT, end_color=TCM_LIGHT, fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
SUB_FONT = Font(name="Arial", bold=True, color=TCM_GREEN, size=10)
BODY_FONT = Font(name="Arial", size=10)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="BDBDBD")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header_row(ws, row: int, max_col: int):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def _style_data_rows(ws, start_row: int, end_row: int, max_col: int):
    for r in range(start_row, end_row + 1):
        fill = PatternFill(start_color="F8F9FA" if r % 2 == 0 else "FFFFFF", end_color="F8F9FA" if r % 2 == 0 else "FFFFFF", fill_type="solid")
        for col in range(1, max_col + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = fill
            cell.font = BODY_FONT
            cell.alignment = LEFT
            cell.border = THIN_BORDER


def _write_df_to_sheet(ws, df: pd.DataFrame, start_row: int = 1, title: Optional[str] = None):
    """Write a DataFrame to a worksheet with formatting."""
    if df is None or df.empty:
        ws.cell(row=start_row, column=1).value = "暂无数据"
        return start_row + 1

    current_row = start_row

    if title:
        ws.cell(row=current_row, column=1).value = title
        ws.cell(row=current_row, column=1).font = Font(name="Arial", bold=True, size=13, color=TCM_GREEN)
        ws.cell(row=current_row, column=1).alignment = LEFT
        current_row += 1

    # Header
    for col_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=current_row, column=col_idx).value = col_name

    _style_header_row(ws, current_row, len(df.columns))
    header_row = current_row
    current_row += 1

    # Data
    for _, row_data in df.iterrows():
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = value
        current_row += 1

    _style_data_rows(ws, header_row + 1, current_row - 1, len(df.columns))

    # Auto-width
    for col_idx, col_name in enumerate(df.columns, 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(col_name)), df.iloc[:, col_idx - 1].astype(str).str.len().max() if not df.empty else 10)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    return current_row


def _add_matplotlib_image(ws, fig: plt.Figure, anchor: str = "A1", width_cm: float = 18):
    """Insert a matplotlib figure as image into worksheet."""
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=120, bbox_inches="tight")
    buf.seek(0)
    plt.close(fig)

    img = Image(buf)
    # Scale image
    pixels_per_cm = 37.8
    img.width = int(width_cm * pixels_per_cm)
    img.height = int(img.width * 0.65)
    ws.add_image(img, anchor)


def generate_excel_report(
    output_path: str,
    herb_names: list,
    disease_name: str,
    compounds_df: pd.DataFrame,
    drug_targets_df: pd.DataFrame,
    disease_targets_df: pd.DataFrame,
    intersection_genes: list,
    ppi_df: pd.DataFrame,
    centrality_df: pd.DataFrame,
    enrichment_results: dict,
    figures: dict,
) -> str:
    """
    Generate comprehensive Excel report with all analysis results.
    Returns the path to the saved file.
    """
    wb = openpyxl.Workbook()

    # ── Sheet 0: Summary ──────────────────────────────────────────────────────
    ws0 = wb.active
    ws0.title = "分析概要"
    ws0.column_dimensions["A"].width = 25
    ws0.column_dimensions["B"].width = 45

    ws0.merge_cells("A1:B1")
    ws0["A1"] = "网络药理学分析报告"
    ws0["A1"].font = Font(name="Arial", bold=True, size=18, color=TCM_GREEN)
    ws0["A1"].alignment = CENTER
    ws0.row_dimensions[1].height = 40

    info = [
        ("分析时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("中药/复方", "、".join(herb_names)),
        ("疾病", disease_name),
        ("活性成分数量", len(compounds_df) if compounds_df is not None else 0),
        ("药物靶点数量", drug_targets_df["Gene"].nunique() if drug_targets_df is not None and "Gene" in drug_targets_df.columns else 0),
        ("疾病靶点数量", len(disease_targets_df) if disease_targets_df is not None else 0),
        ("交集靶点数量", len(intersection_genes)),
        ("KEGG通路数量 (P≤0.05)", len(enrichment_results.get("KEGG", pd.DataFrame())) if enrichment_results else 0),
    ]

    for i, (label, value) in enumerate(info, 3):
        ws0.cell(row=i, column=1).value = label
        ws0.cell(row=i, column=1).font = SUB_FONT
        ws0.cell(row=i, column=1).fill = SUB_FILL
        ws0.cell(row=i, column=2).value = str(value)
        ws0.cell(row=i, column=2).font = BODY_FONT
        for col in [1, 2]:
            ws0.cell(row=i, column=col).border = THIN_BORDER
            ws0.cell(row=i, column=col).alignment = LEFT

    # ── Sheet 1: Active Compounds ─────────────────────────────────────────────
    ws1 = wb.create_sheet("活性成分")
    _write_df_to_sheet(ws1, compounds_df, title="活性成分列表 (ADME筛选后)")

    # ── Sheet 2: Drug Targets ────────────────────────────────────────────────
    ws2 = wb.create_sheet("药物靶点")
    _write_df_to_sheet(ws2, drug_targets_df, title="药物预测靶点")

    # ── Sheet 3: Disease Targets ─────────────────────────────────────────────
    ws3 = wb.create_sheet("疾病靶点")
    _write_df_to_sheet(ws3, disease_targets_df, title=f"疾病相关靶点 ({disease_name})")

    # ── Sheet 4: Intersection ────────────────────────────────────────────────
    ws4 = wb.create_sheet("交集靶点")
    inter_df = pd.DataFrame({"交集靶点": sorted(intersection_genes)})
    _write_df_to_sheet(ws4, inter_df, title="药物-疾病交集靶点")

    # Venn figure
    if "venn" in figures:
        ws4.cell(row=len(intersection_genes) + 5, column=1).value = "韦恩图"
        ws4.cell(row=len(intersection_genes) + 5, column=1).font = SUB_FONT
        _add_matplotlib_image(ws4, figures["venn"], anchor=f"A{len(intersection_genes) + 6}")

    # ── Sheet 5: PPI Network ─────────────────────────────────────────────────
    ws5 = wb.create_sheet("PPI网络")
    _write_df_to_sheet(ws5, centrality_df, title="核心靶点 (按Hub分数排序)")

    if "ppi" in figures:
        row_offset = len(centrality_df) + 5 if centrality_df is not None else 5
        _add_matplotlib_image(ws5, figures["ppi"], anchor=f"A{row_offset}")

    # ── Sheet 6: KEGG ────────────────────────────────────────────────────────
    ws6 = wb.create_sheet("KEGG通路")
    kegg_df = enrichment_results.get("KEGG", pd.DataFrame())
    _write_df_to_sheet(ws6, kegg_df, title="KEGG通路富集分析 (P≤0.05)")

    if "kegg" in figures:
        row_offset = len(kegg_df) + 5 if kegg_df is not None and not kegg_df.empty else 5
        _add_matplotlib_image(ws6, figures["kegg"], anchor=f"A{row_offset}")

    # ── Sheet 7: GO Analysis ─────────────────────────────────────────────────
    ws7 = wb.create_sheet("GO富集分析")
    row = 1
    for go_cat in ["GO_BP", "GO_CC", "GO_MF"]:
        go_df = enrichment_results.get(go_cat, pd.DataFrame())
        row = _write_df_to_sheet(ws7, go_df, start_row=row, title=f"{go_cat} 富集结果") + 2

        cat_key = go_cat.lower()
        if cat_key in figures:
            _add_matplotlib_image(ws7, figures[cat_key], anchor=f"J{row - len(go_df) - 3 if go_df is not None and not go_df.empty else 1}")

    wb.save(output_path)
    return output_path
